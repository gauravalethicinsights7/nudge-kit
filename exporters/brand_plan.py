"""specs/m5-brand-plan.md: 'BrandPlan -> Markdown and DOCX (python-docx) via
exporters/brand_plan.py.' Purely templated rendering of an already-validated
BrandPlan — no LLM involved.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from schemas.brand_plan import BrandPlan


def to_markdown(plan: BrandPlan) -> str:
    lines: list[str] = [f"# Brand Plan ({plan.plan_horizon_months}-month horizon)", ""]

    lines += ["## Situation", plan.situation.summary, ""]
    if plan.situation.key_facts:
        lines += [f"- {f.text}" for f in plan.situation.key_facts] + [""]

    lines += ["## Key Issues", ""]
    for ki in plan.key_issues:
        lines.append(
            f"- **{ki.statement}** — barrier: {ki.barrier}; revenue at stake: "
            f"{ki.revenue_at_stake.value:,.0f} {ki.revenue_at_stake.currency}"
        )
    lines.append("")

    lines += ["## Imperatives", ""]
    for imp in plan.imperatives:
        lines.append(f"- **{imp.title}** ({imp.from_rung.value} -> {imp.to_rung.value})")
    lines.append("")

    lines += [
        "## Positioning",
        f"- Target: {plan.positioning.target}",
        f"- Frame of reference: {plan.positioning.frame_of_reference}",
        f"- Point of difference ({plan.positioning.driver.value}): {plan.positioning.point_of_difference}",
        "",
    ]

    lines += ["## Message House", f"Core: {plan.message_house.core}", ""]
    for pillar in plan.message_house.pillars:
        lines.append(f"- **{pillar.driver.value}**: {pillar.message}")
    lines.append("")

    lines += ["## Objectives", ""]
    for o in plan.objectives:
        lines.append(f"- {o.metric}: {o.baseline:,.0f} -> {o.target:,.0f} by {o.due}")
    lines.append("")

    lines += ["## KPI Tree", "Leading: " + ", ".join(k.name for k in plan.kpi_tree.leading), ""]
    lines.append("Lagging: " + ", ".join(k.name for k in plan.kpi_tree.lagging))
    lines.append("")

    lines += ["## Forecast", ""]
    for scenario_name in ("base", "upside", "downside"):
        scenario = getattr(plan.forecast, scenario_name)
        roi_text = f"{scenario.roi:.2f}" if scenario.roi is not None else "n/a (incremental spend unknown)"
        lines.append(
            f"- **{scenario_name}**: delta_nrx={scenario.delta_nrx:,.1f}, "
            f"revenue={scenario.revenue.value:,.0f} {scenario.revenue.currency}, roi={roi_text}"
        )
    lines.append("")
    lines += ["Key assumptions (base scenario):", ""]
    for a in plan.forecast.base.assumptions:
        lines.append(f"- {a.name} = {a.value:.3f} (confidence {a.confidence:.2f})")
    lines.append("")

    lines += ["## Risks & Tests", ""]
    for r in plan.risks_tests:
        lines.append(f"- {r.assumption} (confidence {r.confidence:.2f}): {r.test_design} -> {r.stop_go}")
    lines.append("")

    lines += ["## Compliance Flags", ""]
    if plan.compliance_flags:
        for f in plan.compliance_flags:
            lines.append(f"- [{f.severity.upper()}] {f.rule_id}: {f.item}")
    else:
        lines.append("None raised by the automated guardrail.")
    lines.append("")

    return "\n".join(lines)


def to_docx(plan: BrandPlan, path: str | Path) -> Path:
    path = Path(path)
    doc = Document()
    doc.add_heading(f"Brand Plan ({plan.plan_horizon_months}-month horizon)", level=0)

    doc.add_heading("Situation", level=1)
    doc.add_paragraph(plan.situation.summary)
    for f in plan.situation.key_facts:
        doc.add_paragraph(f.text, style="List Bullet")

    doc.add_heading("Key Issues", level=1)
    for ki in plan.key_issues:
        doc.add_paragraph(
            f"{ki.statement} — barrier: {ki.barrier}; revenue at stake: "
            f"{ki.revenue_at_stake.value:,.0f} {ki.revenue_at_stake.currency}",
            style="List Bullet",
        )

    doc.add_heading("Imperatives", level=1)
    for imp in plan.imperatives:
        doc.add_paragraph(f"{imp.title} ({imp.from_rung.value} -> {imp.to_rung.value})", style="List Bullet")

    doc.add_heading("Positioning", level=1)
    doc.add_paragraph(f"Target: {plan.positioning.target}")
    doc.add_paragraph(f"Frame of reference: {plan.positioning.frame_of_reference}")
    doc.add_paragraph(
        f"Point of difference ({plan.positioning.driver.value}): {plan.positioning.point_of_difference}"
    )

    doc.add_heading("Message House", level=1)
    doc.add_paragraph(f"Core: {plan.message_house.core}")
    for pillar in plan.message_house.pillars:
        doc.add_paragraph(f"{pillar.driver.value}: {pillar.message}", style="List Bullet")

    doc.add_heading("Objectives", level=1)
    for o in plan.objectives:
        doc.add_paragraph(f"{o.metric}: {o.baseline:,.0f} -> {o.target:,.0f} by {o.due}", style="List Bullet")

    doc.add_heading("Forecast", level=1)
    for scenario_name in ("base", "upside", "downside"):
        scenario = getattr(plan.forecast, scenario_name)
        roi_text = f"{scenario.roi:.2f}" if scenario.roi is not None else "n/a (incremental spend unknown)"
        doc.add_paragraph(
            f"{scenario_name}: delta_nrx={scenario.delta_nrx:,.1f}, "
            f"revenue={scenario.revenue.value:,.0f} {scenario.revenue.currency}, roi={roi_text}",
            style="List Bullet",
        )

    doc.add_heading("Risks & Tests", level=1)
    for r in plan.risks_tests:
        doc.add_paragraph(f"{r.assumption} (confidence {r.confidence:.2f}): {r.test_design} -> {r.stop_go}", style="List Bullet")

    doc.add_heading("Compliance Flags", level=1)
    if plan.compliance_flags:
        for f in plan.compliance_flags:
            doc.add_paragraph(f"[{f.severity.upper()}] {f.rule_id}: {f.item}", style="List Bullet")
    else:
        doc.add_paragraph("None raised by the automated guardrail.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
